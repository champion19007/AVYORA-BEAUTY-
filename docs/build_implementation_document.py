from pathlib import Path
import re, math, textwrap, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
VIS = ROOT / 'implementation-visuals'
VIS.mkdir(exist_ok=True)
INK='#171717'; PURPLE='#EFE3FA'; ACCENT='#DAC2F0'; GRAY='#646464'; LINE='#D9D9D9'; WHITE='#FFFFFF'
def font(size=28,bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf',size)
def canvas(w=1600,h=1000):
    im=Image.new('RGB',(w,h),WHITE); return im,ImageDraw.Draw(im)
def label(d,xy,text,size=28,bold=False,fill=INK,width=None):
    x,y=xy; f=font(size,bold)
    lines=[]
    for raw in text.split('\n'):
        if not width: lines.append(raw); continue
        line=''
        for word in raw.split():
            candidate=(line+' '+word).strip()
            if d.textlength(candidate,font=f)>width and line: lines.append(line); line=word
            else: line=candidate
        lines.append(line)
    for line in lines:
        d.text((x,y),line,font=f,fill=fill); y+=int(size*1.35)
    return y
def box(d,rect,title,body='',fill=PURPLE,size=28):
    d.rounded_rectangle(rect,radius=18,fill=fill,outline=LINE,width=2)
    x,y,x2,y2=rect; yy=label(d,(x+20,y+17),title,size,True,width=x2-x-40)
    if body: label(d,(x+20,yy+12),body,size-4,width=x2-x-40)
def arrow(d,a,b,fill=GRAY):
    d.line([a,b],fill=fill,width=4)
    angle=math.atan2(b[1]-a[1],b[0]-a[0]); r=14
    pts=[b,(b[0]-r*math.cos(angle-.5),b[1]-r*math.sin(angle-.5)),(b[0]-r*math.cos(angle+.5),b[1]-r*math.sin(angle+.5))]
    d.polygon(pts,fill=fill)
def save(im,name): im.save(VIS/name)

im,d=canvas(1600,1040)
label(d,(55,30),'Cached storefront with client inference',45,True)
box(d,(50,115,510,300),'Public CDN','HTML, portraits, product assets\nVersioned catalogue and knowledge',WHITE)
box(d,(585,115,1050,300),'Customer browser','Quiz and Bayesian inference\nOptional local cosmetic model')
box(d,(1130,115,1550,300),'Client screens','Routine, bag, account\nFresh commerce quote',WHITE)
arrow(d,(510,210),(585,210)); arrow(d,(1050,210),(1130,210))
box(d,(585,400,1050,590),'Authorized Next.js API','Identity, consent, saved routines\nPricing, stock, orders and payments')
arrow(d,(820,300),(820,400)); arrow(d,(1340,300),(1050,470))
box(d,(50,700,475,905),'PostgreSQL','Existing commerce and job queue\nProfiles, releases and decision traces',WHITE)
box(d,(570,700,1010,905),'Private scan storage','Short photo lifetime\nOwned access and deletion',WHITE)
box(d,(1095,700,1550,905),'Optional hosted worker','Evaluated image inference\nBounded queue, quota and spend',WHITE)
arrow(d,(690,590),(265,700)); arrow(d,(820,590),(790,700)); arrow(d,(1050,540),(1320,700))
label(d,(55,958),'Core rules run in both browser and server. Money and permissions remain server authoritative.',27,width=1470)
save(im,'architecture.png')

im,d=canvas(1600,1080)
label(d,(50,30),'Data relationships and immutable releases',44,True)
box(d,(50,115,480,270),'users','Existing identity primary key',WHITE)
box(d,(580,115,1010,270),'catalog_products','Variants and formulation versions',WHITE)
box(d,(1110,115,1550,270),'ingredients','Canonical IDs and reviewed cautions',WHITE)
box(d,(50,385,480,590),'skin_profiles','One owner, versioned answers\nConsent and expiration')
box(d,(580,385,1010,590),'formulations','Full INCI and product ingredients\nApproved usage profiles')
box(d,(1110,385,1550,590),'knowledge records','Rules, interactions, evidence\nApproved explanations')
box(d,(50,725,480,930),'routine_results','Versioned plan and trace\nNormalized routine_days',WHITE)
box(d,(580,725,1010,930),'scan_sessions','Owned result and consent\nPrivate photo key and expiry',WHITE)
box(d,(1110,725,1550,930),'kb_releases','Published immutable JSON\nBayesian and model versions',WHITE)
arrow(d,(265,270),(265,385)); arrow(d,(265,590),(265,725)); arrow(d,(480,835),(580,835)); arrow(d,(795,270),(795,385)); arrow(d,(1010,485),(1110,485)); arrow(d,(1330,270),(1330,385)); arrow(d,(1330,590),(1330,725)); arrow(d,(1110,865),(480,865))
label(d,(50,986),'Orders retain historical charged values. Personal records never enter the public knowledge release.',27,width=1490)
save(im,'data-model.png')

im,d=canvas(1600,990)
label(d,(55,35),'A small Bayesian update',44,True)
label(d,(55,105),'Illustrative arithmetic only. Production likelihoods require evaluated data.',28,fill=GRAY)
values=[.2,.428571,.6]; labels=['Prior','Accepted evidence A','Independent evidence B']
for i,(v,t) in enumerate(zip(values,labels)):
    x=160+i*480; y=775-int(v*650)
    d.rounded_rectangle((x,y,x+270,775),radius=14,fill=ACCENT)
    label(d,(x+42,y-60),f'{v*100:.1f}%',40,True)
    label(d,(x-15,810),t,26,width=380)
d.line((90,775,1500,775),fill=GRAY,width=3)
for i in range(1,5):
    yy=775-i*130; d.line((90,yy,1500,yy),fill='#EEEEEE',width=2); label(d,(35,yy-14),str(i*20),20,fill=GRAY)
box(d,(130,185,1480,315),'Prior odds 0.25  →  × 3 = 0.75  →  × 2 = 1.50','Convert odds back using odds / (1 + odds). Missing evidence has ratio 1.',WHITE,28)
label(d,(55,905),'Correlated quiz and photo observations share one evidence group and must not be counted twice.',28,width=1480)
save(im,'bayesian-update.png')

im,d=canvas(1600,1110)
label(d,(50,30),'Routine generation with hard constraints',43,True)
stages=[('1  Validate profile','Typed answers, owned products, budget and unknown states'),('2  Resolve knowledge','Canonical ingredients, formulation coverage and approved directions'),('3  Apply exclusions','Safety policy, irritation, age, prescriptions and allergy'),('4  Combine evidence','Bounded Bayesian updates; grouped correlated observations'),('5  Select candidates','Stock, affordability, step cap and suitability ranking'),('6  Build weekly sessions','Actual AM and PM day slots with frequency and ordering'),('7  Revalidate and explain','Whole-plan conflicts, optional slots, trace and fresh quote')]
for i,(t,b) in enumerate(stages):
    yy=110+i*133; box(d,(150,yy,1450,yy+106),t,b,WHITE if i%2 else PURPLE,28)
    if i<6: arrow(d,(800,yy+106),(800,yy+133))
label(d,(55,1050),'A score never overrides an exclusion. An empty optional treatment is a valid outcome.',27,width=1470)
save(im,'routine-pipeline.png')

im,d=canvas(1600,1120)
label(d,(45,30),'Optional scan paths and photo boundaries',44,True)
box(d,(440,115,1160,250),'Capture after explanation and consent','One face, supported lighting, quality accepted',PURPLE,29)
box(d,(85,345,710,535),'Local inference','Validated browser model\nPhoto remains in device memory',WHITE,30)
box(d,(885,345,1510,535),'Hosted inference','Owned private upload and durable quota\nWorker rechecks consent and runs model',WHITE,30)
arrow(d,(650,250),(395,345)); arrow(d,(960,250),(1195,345))
label(d,(675,285),'OR',27,True)
box(d,(85,610,710,750),'Local observations','Clear photo buffers and camera tracks',PURPLE,28)
box(d,(885,610,1510,750),'Owned hosted observations','Delete raw photo after processing',PURPLE,28)
arrow(d,(395,535),(395,610)); arrow(d,(1195,535),(1195,610))
box(d,(350,860,1250,990),'Combine evidence and build a routine','Hard rules, Bayesian fusion and actual weekly sessions',WHITE,30)
arrow(d,(395,750),(590,860)); arrow(d,(1195,750),(1010,860))
label(d,(45,1040),'Either path can decline an image. The quiz remains available without any photo.',27,width=1490)
save(im,'scan-sequence.png')

im,d=canvas(1600,1280)
label(d,(45,25),'Desktop wireframe following the current live Nuve site',39,True)
d.rectangle((40,105,1560,890),fill='#87949F')
d.ellipse((470,120,1040,885),fill='#75818C')
label(d,(560,440),'Licensed full-bleed\nportrait photograph',28,fill=WHITE,width=380)
label(d,(80,135),'AVYORA',36,fill=WHITE)
for yy in [148,158,168]: d.line((1480,yy,1518,yy),fill=WHITE,width=3)
label(d,(1110,235),'A routine for your skin,\nbudget and tolerance.',28,fill=WHITE,width=370)
d.rounded_rectangle((1220,380,1510,450),radius=35,fill=WHITE)
label(d,(1250,400),'Find My Routine',26)
label(d,(80,440),'CARE WITH\nCLARITY',25,fill=WHITE)
label(d,(80,655),'Skin care for\nyour daily routine',82,fill=WHITE)
label(d,(1260,845),'Shop in the menu',20,fill=WHITE)
label(d,(45,935),'Keep the live reference section sequence',32,True)
box(d,(45,1005,770,1190),'Editorial and shop sections','About  →  Results grid  →  Vision\nSix Features  →  Services  →  Genuine reviews',WHITE,28)
box(d,(815,1005,1555,1190),'Routine and support sections','Process and routine card  →  Full-width image\nFAQ  →  Support form  →  Footer',WHITE,28)
label(d,(45,1220),'Illustrative wireframe. Exact photography, font rendering and motion are checked against live captures.',23,fill=GRAY,width=1500)
save(im,'storefront-wireframes.png')

im,d=canvas(1600,1130)
label(d,(50,25),'Proposed routine result screen',42,True)
label(d,(50,100),'Your essential routine',38,True)
label(d,(50,160),'Based on your answers. Photo observations are optional and may be uncertain.',25,width=1450)
box(d,(50,225,1060,345),'Morning','Cleanse if needed  →  Moisturize  →  Daytime sunscreen','#F1F1F1',30)
box(d,(50,375,1060,495),'Evening','Cleanse  →  Moisturize  ·  Elective treatment only when eligible',WHITE,30)
box(d,(1130,225,1540,495),'New spend','Owned items excluded\nVariant-specific quote\nOptional additions separate',WHITE,29)
label(d,(50,550),'Your week',32,True)
days=['Mon','Tue','Wed','Thu','Fri','Sat','Sun']
for i,t in enumerate(days):
    x=50+i*215; box(d,(x,615,x+195,815),t,'AM essentials\nPM essentials','#F1F1F1' if i%2 else WHITE,24)
box(d,(50,865,755,1030),'Why this plan','Fit your budget and daily step cap.\nNo suitable elective treatment is also a valid result.',WHITE,27)
box(d,(800,865,1540,1030),'Edit and save','Edit answers  ·  Check substitutions\nSave privately  ·  Give weekly feedback',WHITE,27)
save(im,'routine-screen.png')

def shade(cell,color):
    pr=cell._tc.get_or_add_tcPr(); el=OxmlElement('w:shd'); el.set(qn('w:fill'),color); pr.append(el)
def hyperlink(p,text,url):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'235B91'); pr.append(c); r.append(pr)
    t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)
def text_into(p,s):
    parts=re.split(r'(https?://[^\s]+)',s)
    for part in parts:
        if part.startswith('http'): hyperlink(p,part,part)
        elif part: p.add_run(part.replace('**','').replace('`',''))
def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    available=6.9
    widths=([1.55,2.65,2.70] if len(headers)==3 else [1.55,1.6,1.9,1.85])
    if len(headers)==2: widths=[2.1,4.8]
    if len(headers)>4: widths=[available/len(headers)]*len(headers)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    props.append(borders)
    for i,h in enumerate(headers):
        t.columns[i].width=Inches(widths[i]); cell=t.rows[0].cells[i]; cell.width=Inches(widths[i]); shade(cell,'E8E8E8'); text_into(cell.paragraphs[0],h)
        for r in cell.paragraphs[0].runs: r.bold=True
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for i,s in enumerate(values):
            cells[i].width=Inches(widths[i]); text_into(cells[i].paragraphs[0],s)
    for row in t.rows:
        cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'95'); e.set(qn('w:type'),'dxa'); mar.append(e)
            c._tc.get_or_add_tcPr().append(mar)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1.05
                for r in p.runs: r.font.size=Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=Inches(.7); section.bottom_margin=Inches(.7); section.left_margin=Inches(.8); section.right_margin=Inches(.8)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[n]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(29); doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(19); doc.styles['Heading 1'].font.bold=True; doc.styles['Heading 1'].paragraph_format.space_before=Pt(17); doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 2'].font.size=Pt(14)
code=doc.styles.add_style('Code',1); code.font.name='Consolas'; code.font.size=Pt(9); code.paragraph_format.space_after=Pt(0); code.paragraph_format.line_spacing=1
for s in [doc.styles['Title'],doc.styles['Subtitle'],doc.styles['Heading 1'],doc.styles['Heading 2']]:
    pr=s.element.find(qn('w:pPr'))
    if pr is not None:
        for el in list(pr):
            if el.tag==qn('w:pBdr'): pr.remove(el)

footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Avyora developer specification  |  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='Avyora Developer Implementation Specification'; doc.core_properties.author='Avyora'; doc.core_properties.subject='Storefront redesign and knowledge based personalization'

lines=(ROOT/'avyora-developer-implementation-specification-2026-10-07.md').read_text(encoding='utf-8').splitlines()
i=0; incode=False; images=0
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        incode=not incode; i+=1; continue
    if incode:
        p=doc.add_paragraph(style='Code'); p.add_run(line); p.paragraph_format.keep_together=True; i+=1; continue
    if not line.strip(): i+=1; continue
    if line.startswith('# '): doc.add_paragraph(line[2:],style='Title'); i+=1; continue
    if line.startswith('## '):
        heading=line[3:]
        # Start each major chapter on a fresh page for reliable navigation.
        doc.add_page_break(); doc.add_heading(heading,level=1); i+=1; continue
    if line.startswith('!['):
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if m:
            p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
            run=p.add_run(); pic=run.add_picture(str(ROOT/m.group(2)),width=Inches(6.9)); pic._inline.docPr.set('descr',m.group(1)); images+=1
            cp=doc.add_paragraph(m.group(1)); cp.style=doc.styles['Caption']; cp.paragraph_format.space_after=Pt(10)
        i+=1; continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].startswith('|'):
            block.append(lines[i]); i+=1
        headers=[x.strip() for x in block[0].strip('|').split('|')]
        rows=[[x.strip() for x in row.strip('|').split('|')] for row in block[2:]]
        # All document tables have simple pipe-free values.
        table(headers,rows); continue
    p=doc.add_paragraph(); text_into(p,line)
    i+=1

output=ROOT/'Avyora Developer Implementation Specification.docx'
doc.save(output)
print(json.dumps({'docx':str(output),'images':images,'tables':len(doc.tables),'paragraphs':len(doc.paragraphs),'words':len(' '.join(lines).split())}))
